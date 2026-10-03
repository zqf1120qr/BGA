# -*- coding: utf-8 -*-
"""
🌟 BGA 质检数据 Excel 结构化报表生成器 (Excel Reporter)
-------------------------------------------------------
将整板质检汇总指标及所有焊点的几何坐标、气泡率、虚焊比、短路状态导出为美观的 .xlsx 表格。
"""
from __future__ import annotations

import datetime
from typing import Any, Dict, List
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from desktop_app.core.inspection_manager import SolderBallRecord


def export_inspection_excel(
    save_path: str,
    image_name: str,
    summary_stats: Dict[str, Any],
    records: List[SolderBallRecord],
):
    wb = openpyxl.Workbook()

    # --- Sheet 1: 质检汇总总览 ---
    ws_summary = wb.active
    ws_summary.title = "质检汇总报告"
    ws_summary.views.sheetView[0].showGridLines = True

    # 样式定义
    font_title = Font(name="微软雅黑", size=16, bold=True, color="1E3A8A")
    font_bold = Font(name="微软雅黑", size=11, bold=True)
    font_norm = Font(name="微软雅黑", size=11)
    
    fill_pass = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    fill_ng = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    fill_header = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    font_header = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    ws_summary["A1"] = "🌟 BGA 智能工业质检分析报表"
    ws_summary["A1"].font = font_title
    ws_summary.merge_cells("A1:D1")

    now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    summary_rows = [
        ("检测图片名称", image_name),
        ("报告导出时间", now_str),
        ("整板最终判定", summary_stats.get("board_status", "UNKNOWN")),
        ("焊球总数 (Total)", summary_stats.get("total_solders", 0)),
        ("气泡超标缺陷数", summary_stats.get("void_ng_count", 0)),
        ("桥连短路缺陷数", summary_stats.get("bridge_count", 0)),
        ("虚焊少锡缺陷数", summary_stats.get("insufficient_count", 0)),
        ("最大气泡空洞率", f"{summary_stats.get('max_void_rate', 0.0) * 100:.2f}%"),
        ("最大面积偏离度", f"-{summary_stats.get('max_reduction_percent', 0.0):.2f}%"),
        ("质检耗时 (ms)", f"{summary_stats.get('elapsed_ms', 0.0):.1f}"),
    ]

    for idx, (k, v) in enumerate(summary_rows, start=3):
        ws_summary.cell(row=idx, column=1, value=k).font = font_bold
        ws_summary.cell(row=idx, column=1).border = thin_border
        
        val_cell = ws_summary.cell(row=idx, column=2, value=v)
        val_cell.font = font_norm
        val_cell.border = thin_border

        if k == "整板最终判定":
            val_cell.font = Font(name="微软雅黑", size=12, bold=True, color="15803D" if v == "PASS" else "B91C1C")
            val_cell.fill = fill_pass if v == "PASS" else fill_ng

    # --- Sheet 2: 焊点全量明细表 ---
    ws_details = wb.create_sheet(title="焊点全量明细")
    ws_details.views.sheetView[0].showGridLines = True

    headers = [
        "焊球序号", "中心X坐标(px)", "中心Y坐标(px)", "实测半径(px)",
        "气泡占比(%)", "少锡面积偏离(%)", "桥连短路状态", "检出置信度(%)", "单点结论"
    ]
    ws_details.append(headers)

    for col_idx in range(1, len(headers) + 1):
        cell = ws_details.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    for rec in records:
        row_vals = [
            f"#{rec.index}",
            round(rec.cx, 1),
            round(rec.cy, 1),
            round(rec.radius, 1),
            round(rec.void_rate * 100, 2),
            round(-rec.reduction_percent, 2) if rec.reduction_percent > 0 else 0.0,
            "⚠️短路" if rec.is_bridge else "正常",
            round(rec.confidence * 100, 1),
            rec.status,
        ]
        ws_details.append(row_vals)
        cur_row = ws_details.max_row
        
        for col_idx in range(1, len(row_vals) + 1):
            c = ws_details.cell(row=cur_row, column=col_idx)
            c.font = font_norm
            c.border = thin_border
            c.alignment = Alignment(horizontal="center", vertical="center")
            if rec.is_ng:
                c.fill = fill_ng
            else:
                c.fill = fill_pass

    # 自适应列宽
    for ws in (ws_summary, ws_details):
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(save_path)
