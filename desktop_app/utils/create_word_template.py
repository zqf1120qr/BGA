# -*- coding: utf-8 -*-
"""
生成工业级 X射线 质检报告 Word 模板 (.docx)
1:1 完美复刻现有 PDF 报告版式与视觉规范，提供清晰的占位符变量供用户可视化修改排版。
"""
from __future__ import annotations

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn


def set_cell_background(cell, hex_color: str):
    """设置单元格背景颜色 (hex 字符串如 '1E293B')"""
    hex_color = hex_color.lstrip('#')
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    """设置单元格内边距 (单位 dxa, 20 dxa = 1 pt)"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    """为表格设置精致细腻的边框"""
    tblPr = table._tbl.tblPr
    borders_elm = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders_elm)


def create_bga_report_template(output_path: str):
    doc = Document()

    # 1. 设置页面边距为 1.5 cm (与 PDF 完全对齐)
    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)
        section.page_width = Cm(21.0)   # A4
        section.page_height = Cm(29.7)

    # 默认正文字体
    normal_style = doc.styles['Normal']
    normal_style.font.name = '微软雅黑'
    normal_style._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    normal_style.font.size = Pt(9.5)
    normal_style.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    # 2. 报告主标题
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(8)
    run_title = p_title.add_run("X射线 工业检测分析报告单")
    run_title.font.name = '微软雅黑'
    run_title._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)  # 工业深蓝

    # 3. 基础信息与整板大判定栏 (表格 2 行 3 列)
    tbl_info = doc.add_table(rows=2, cols=3)
    tbl_info.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_info, color="CBD5E1")

    # 列宽设置 (合计约 18 cm)
    col_widths_info = [Cm(6.5), Cm(6.8), Cm(4.7)]
    for row in tbl_info.rows:
        for idx, width in enumerate(col_widths_info):
            row.cells[idx].width = width

    # 第一行
    # (0, 0)
    c00 = tbl_info.cell(0, 0)
    set_cell_margins(c00, top=140, bottom=140, left=160, right=160)
    p = c00.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run("待测图片: ")
    r1.font.bold = True
    p.add_run("{{IMAGE_NAME}}")

    # (0, 1)
    c01 = tbl_info.cell(0, 1)
    set_cell_margins(c01, top=140, bottom=140, left=160, right=160)
    p = c01.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run("检测模式: ")
    r1.font.bold = True
    p.add_run("{{INSPECT_MODE}}")

    # (0, 2) 判定大标签表头
    c02 = tbl_info.cell(0, 2)
    set_cell_margins(c02, top=140, bottom=80, left=160, right=160)
    set_cell_background(c02, "DCFCE7")  # 默认浅绿色背景
    p = c02.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run("整板最终判定")
    r1.font.bold = True
    r1.font.size = Pt(10)
    r1.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    # 第二行
    # (1, 0)
    c10 = tbl_info.cell(1, 0)
    set_cell_margins(c10, top=140, bottom=140, left=160, right=160)
    p = c10.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run("检验时间: ")
    r1.font.bold = True
    p.add_run("{{INSPECT_TIME}}")

    # (1, 1)
    c11 = tbl_info.cell(1, 1)
    set_cell_margins(c11, top=140, bottom=140, left=160, right=160)
    p = c11.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run("检测用时: ")
    r1.font.bold = True
    p.add_run("{{ELAPSED_MS}} ms")

    # (1, 2) 判定大印章结果
    c12 = tbl_info.cell(1, 2)
    set_cell_margins(c12, top=80, bottom=140, left=160, right=160)
    set_cell_background(c12, "DCFCE7")  # 默认浅绿色背景
    p = c12.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r_stamp = p.add_run("{{BOARD_STATUS}}")
    r_stamp.font.size = Pt(15)
    r_stamp.font.bold = True
    r_stamp.font.color.rgb = RGBColor(0x16, 0xA3, 0x4A)

    # 4. 小节标题：多维指标统计
    p_sec1 = doc.add_paragraph()
    p_sec1.paragraph_format.space_before = Pt(8)
    p_sec1.paragraph_format.space_after = Pt(3)
    r = p_sec1.add_run("【多维指标统计】")
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    # 指标统计表格 (2 行 6 列)
    tbl_stats = doc.add_table(rows=2, cols=6)
    tbl_stats.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_stats, color="CBD5E1")

    headers_stats = ["检测焊球数", "空洞超标", "桥连短路", "虚焊少锡", "排除误检(复核)", "最大空洞比"]
    placeholders_stats = [
        "{{TOTAL_SOLDERS}}",
        "{{VOID_NG_COUNT}}",
        "{{BRIDGE_COUNT}}",
        "{{INSUFFICIENT_COUNT}}",
        "{{MANUAL_CLEARED_COUNT}}",
        "{{MAX_VOID_RATE}}"
    ]
    col_w_stat = Cm(3.0)

    # 表头行
    for i, h_text in enumerate(headers_stats):
        cell = tbl_stats.cell(0, i)
        cell.width = col_w_stat
        set_cell_background(cell, "1E293B")  # 深色表头
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h_text)
        run.font.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # 数据行
    for i, p_text in enumerate(placeholders_stats):
        cell = tbl_stats.cell(1, i)
        cell.width = col_w_stat
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(p_text)
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    # 5. 小节标题：全景检测标注图
    p_sec2 = doc.add_paragraph()
    p_sec2.paragraph_format.space_before = Pt(8)
    p_sec2.paragraph_format.space_after = Pt(3)
    r = p_sec2.add_run("【全景检测标注图 (仅高亮显示异常与复核部位)】")
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    # 图片占位段落 (居中显示占位标识，生成报告时引擎会自动将此处替换为实际高清标注图)
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(4)
    p_img.paragraph_format.space_after = Pt(4)
    r_img = p_img.add_run("{{PREVIEW_IMAGE}}")
    r_img.font.size = Pt(10)
    r_img.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    r_img.font.italic = True

    # 6. 图例条表格 (1 行 5 列)
    tbl_legend = doc.add_table(rows=1, cols=5)
    tbl_legend.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_legend, color="CBD5E1")

    legends = [
        ("● 空洞超标", RGBColor(0xD9, 0x77, 0x06)),
        ("● 桥连短路", RGBColor(0xDC, 0x26, 0x26)),
        ("● 虚焊少锡", RGBColor(0x93, 0x33, 0xEA)),
        ("● 人工排除", RGBColor(0x08, 0x91, 0xB2)),
        ("○ 正常合格(底图零遮挡)", RGBColor(0x16, 0xA3, 0x4A)),
    ]
    col_w_leg = [Cm(3.4), Cm(3.4), Cm(3.4), Cm(3.4), Cm(4.4)]
    for i, (l_text, l_color) in enumerate(legends):
        cell = tbl_legend.cell(0, i)
        cell.width = col_w_leg[i]
        set_cell_background(cell, "F8FAFC")
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(l_text)
        run.font.size = Pt(8.5)
        run.font.bold = True
        run.font.color.rgb = l_color

    # 7. 小节标题：不良与复核焊点明细清单
    p_sec3 = doc.add_paragraph()
    p_sec3.paragraph_format.space_before = Pt(8)
    p_sec3.paragraph_format.space_after = Pt(3)
    r = p_sec3.add_run("【不良与复核焊点明细清单】")
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    # 缺陷清单表格
    tbl_defect = doc.add_table(rows=2, cols=7)
    tbl_defect.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_defect, color="CBD5E1")

    defect_headers = ["序号", "中心坐标(X, Y)", "空洞占比", "少锡缩减", "桥连短路", "初始原因", "最终结论"]
    defect_col_w = [Cm(1.6), Cm(2.8), Cm(2.2), Cm(2.2), Cm(2.2), Cm(4.0), Cm(3.0)]

    # 表头
    for i, h_text in enumerate(defect_headers):
        cell = tbl_defect.cell(0, i)
        cell.width = defect_col_w[i]
        set_cell_background(cell, "475569")  # 灰蓝色表头
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h_text)
        run.font.bold = True
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # 模板数据占位行 (第 2 行，报告生成引擎会自动根据真实缺陷数据扩展此行)
    sample_defect_row = [
        "{{DEFECT_ROW}}",
        "(X, Y)",
        "0.0%",
        "0.0%",
        "正常/NG",
        "异常说明",
        "最终结论"
    ]
    for i, d_text in enumerate(sample_defect_row):
        cell = tbl_defect.cell(1, i)
        cell.width = defect_col_w[i]
        set_cell_background(cell, "FFF1F2")  # 默认浅红
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(d_text)
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    # 8. 底部签字栏表格 (1 行 3 列)
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(6)
    p_space.paragraph_format.space_after = Pt(0)

    tbl_sign = doc.add_table(rows=1, cols=3)
    tbl_sign.alignment = WD_TABLE_ALIGNMENT.CENTER
    # 签字栏不显示外框线
    for row in tbl_sign.rows:
        for idx, width in enumerate([Cm(6.0), Cm(6.0), Cm(6.0)]):
            row.cells[idx].width = width

    sign_texts = [
        "品质检验员 (Inspector): ________________",
        "人工复核员 (Reviewer): ________________",
        "审核日期 (Date): ____年__月__日"
    ]
    for i, s_text in enumerate(sign_texts):
        cell = tbl_sign.cell(0, i)
        set_cell_margins(cell, top=80, bottom=80, left=60, right=60)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(s_text)
        run.font.size = Pt(9)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    # 保存文档
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc.save(output_path)
    print(f"[OK] Word 模板已成功生成: {output_path}")


if __name__ == "__main__":
    target = os.path.join(os.path.dirname(__file__), "..", "..", "bga_report_template.docx")
    create_bga_report_template(os.path.abspath(target))
