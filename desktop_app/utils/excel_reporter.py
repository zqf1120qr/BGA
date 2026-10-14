# -*- coding: utf-8 -*-
"""
X射线 质检数据 Excel 结构化报表生成器 (Excel Reporter)
-------------------------------------------------------
将整板质检汇总指标及所有焊点的几何坐标、空洞率、虚焊比、短路状态导出为美观的 .xlsx 表格。
支持单张质检报表导出与多工件批量合并报表 (多 Sheet 结构) 导出。
"""
from __future__ import annotations

import datetime
import os
import re
from typing import Any, Dict, List, Optional

from desktop_app.core.inspection_manager import SolderBallRecord


def _sanitize_sheet_title(index: int, raw_name: str, used_titles: set) -> str:
    """生成合规且唯一的 openpyxl 工作表名称 (≤31字符，无非法字符)"""
    base = os.path.splitext(os.path.basename(raw_name))[0]
    # 清理非法字符: \ / ? * [ ] :
    clean_base = re.sub(r'[\\/*?\[\]:]', '_', base).strip()
    if not clean_base:
        clean_base = f"Board_{index}"

    prefix = f"{index:02d}_"
    max_base_len = 31 - len(prefix)
    candidate = f"{prefix}{clean_base[:max_base_len]}"

    counter = 1
    while candidate in used_titles or len(candidate) == 0:
        suffix = f"_{counter}"
        avail_len = 31 - len(prefix) - len(suffix)
        candidate = f"{prefix}{clean_base[:max(1, avail_len)]}{suffix}"
        counter += 1

    used_titles.add(candidate)
    return candidate


# 焊点全量明细表标准列宽映射 (免去全表数万单元格逐字扫描的庞大 CPU 延迟，极大提速多 Sheet 导出)
SOLDER_COL_WIDTHS = {
    "A": 12,  # 焊球序号
    "B": 15,  # 中心X坐标(px)
    "C": 15,  # 中心Y坐标(px)
    "D": 14,  # 实测半径(px)
    "E": 14,  # 空洞占比(%)
    "F": 15,  # 少锡缩减率(%)
    "G": 14,  # 桥连短路状态
    "H": 15,  # 识别可靠度(%)
    "I": 16,  # 单点结论
    "J": 15,  # 人工复核状态
    "K": 28,  # 复核说明与时间
}


def _apply_detail_sheet_col_widths(ws):
    """为包含海量焊球的明细工作表秒级应用预设高精列宽并开启网格线"""
    ws.views.sheetView[0].showGridLines = True
    for col_letter, width in SOLDER_COL_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width


def _apply_sheet_styles_and_col_width(ws, openpyxl_utils_col_letter):
    """自适应列宽并强制开启网格线 (针对总览看板等小行数工作表)"""
    ws.views.sheetView[0].showGridLines = True
    for col in ws.columns:
        max_len = 0
        for cell in col:
            val_str = str(cell.value or '')
            if "\n" in val_str:
                line_lens = [len(l) for l in val_str.split("\n")]
                max_len = max(max_len, max(line_lens) if line_lens else 0)
            else:
                if val_str.startswith("="):
                    max_len = max(max_len, 14)
                else:
                    zh_len = sum(1.5 if ord(c) > 127 else 1.0 for c in val_str)
                    max_len = max(max_len, int(zh_len))
        col_letter = openpyxl_utils_col_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 45)


def _append_solder_records_table(
    ws,
    records: List[SolderBallRecord],
    start_row: int,
    styles: Dict[str, Any]
):
    """向指定工作表中追加焊点全量明细表 (复用预构建样式对象，零重复实例分配)"""
    headers = [
        "焊球序号", "中心X坐标(px)", "中心Y坐标(px)", "实测半径(px)",
        "空洞占比(%)", "少锡缩减率(%)", "桥连短路状态", "识别可靠度(%)",
        "单点结论", "人工复核状态", "复核说明与时间"
    ]

    align_center = styles["align_center"]
    font_header = styles["font_header"]
    fill_header = styles["fill_header"]
    thin_border = styles["thin_border"]
    font_norm = styles["font_norm"]
    fill_rev = styles["fill_rev"]
    fill_ng = styles["fill_ng"]
    fill_pass = styles["fill_pass"]

    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border

    cur_row = start_row
    for rec in records:
        cur_row += 1
        is_rev = (rec.is_manually_reviewed and rec.manual_status == "PASS")
        rev_status = "已排除误检" if is_rev else "--"
        rev_info = f"{rec.review_reason} ({rec.review_timestamp})" if is_rev else "--"

        row_vals = [
            f"#{rec.index}",
            round(rec.cx, 1),
            round(rec.cy, 1),
            round(rec.radius, 1),
            round(rec.void_rate * 100, 2),
            round(-rec.reduction_percent, 2) if rec.reduction_percent > 0 else 0.0,
            "桥连短路NG" if rec.is_bridge else "正常",
            round(rec.confidence * 100, 1),
            rec.final_status,
            rev_status,
            rev_info,
        ]

        # 行级复用填充色
        target_fill = fill_rev if is_rev else (fill_ng if rec.final_is_ng else fill_pass)

        for col_idx, val in enumerate(row_vals, start=1):
            c = ws.cell(row=cur_row, column=col_idx, value=val)
            c.font = font_norm
            c.border = thin_border
            c.alignment = align_center
            c.fill = target_fill


def export_inspection_excel(
    save_path: str,
    image_name: str,
    summary_stats: Dict[str, Any],
    records: List[SolderBallRecord],
    progress_callback: Optional[Any] = None,
):
    """单张图像质检数据导出为 Excel 表格"""
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()

    # 样式配置 (单例预初始化，消除重复实例化开销)
    font_title = Font(name="微软雅黑", size=16, bold=True, color="1E3A8A")
    font_bold = Font(name="微软雅黑", size=11, bold=True)
    font_norm = Font(name="微软雅黑", size=11)
    font_header = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")

    fill_pass = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    fill_ng = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    fill_rev = PatternFill(start_color="CFFAFE", end_color="CFFAFE", fill_type="solid")
    fill_header = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    style_pack = {
        "font_bold": font_bold,
        "font_norm": font_norm,
        "font_header": font_header,
        "fill_pass": fill_pass,
        "fill_ng": fill_ng,
        "fill_rev": fill_rev,
        "fill_header": fill_header,
        "thin_border": thin_border,
        "align_center": align_center,
        "align_left": align_left,
    }

    if progress_callback:
        progress_callback(1, 3, "正在构建质检汇总报告...")

    # --- Sheet 1: 质检汇总报告 ---
    ws_summary = wb.active
    ws_summary.title = "质检汇总报告"

    ws_summary["A1"] = "X射线 工业检测分析报表"
    ws_summary["A1"].font = font_title
    ws_summary.merge_cells("A1:D1")

    now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    summary_rows = [
        ("检测图片名称", image_name),
        ("报告导出时间", now_str),
        ("整板最终判定", summary_stats.get("board_status", "UNKNOWN")),
        ("检测焊球总数", summary_stats.get("total_solders", 0)),
        ("空洞超标焊点数", summary_stats.get("void_ng_count", 0)),
        ("桥连短路对数", summary_stats.get("bridge_count", 0)),
        ("虚焊少锡焊点数", summary_stats.get("insufficient_count", 0)),
        ("人工排除误检数", summary_stats.get("manual_cleared_count", 0)),
        ("最大空洞占比", f"{float(summary_stats.get('max_void_rate') or 0.0) * 100:.2f}%"),
        ("最大少锡缩减率", f"-{float(summary_stats.get('max_reduction_percent') or 0.0):.2f}%"),
        ("检测用时 (ms)", f"{float(summary_stats.get('elapsed_ms') or 0.0):.1f}"),
    ]

    for idx, (k, v) in enumerate(summary_rows, start=3):
        ws_summary.cell(row=idx, column=1, value=k).font = font_bold
        ws_summary.cell(row=idx, column=1).border = thin_border

        val_cell = ws_summary.cell(row=idx, column=2, value=v)
        val_cell.font = font_norm
        val_cell.border = thin_border

        if k == "整板最终判定":
            is_pass = str(v).startswith("PASS")
            val_cell.font = Font(name="微软雅黑", size=12, bold=True, color="15803D" if is_pass else "B91C1C")
            val_cell.fill = fill_pass if is_pass else fill_ng
        elif k == "人工排除误检数" and int(v) > 0:
            val_cell.font = Font(name="微软雅黑", size=11, bold=True, color="0891B2")
            val_cell.fill = fill_rev

    if progress_callback:
        progress_callback(2, 3, "正在构建焊点全量明细表...")

    # --- Sheet 2: 焊点全量明细表 ---
    ws_details = wb.create_sheet(title="焊点全量明细")
    _append_solder_records_table(ws_details, records, start_row=1, styles=style_pack)

    # 自适应列宽与边框网格 (明细表直接采用极速预置列宽，汇总表自适应)
    _apply_sheet_styles_and_col_width(ws_summary, get_column_letter)
    _apply_detail_sheet_col_widths(ws_details)

    if progress_callback:
        progress_callback(3, 3, "正在保存 Excel 表格...")

    wb.save(save_path)


def export_batch_inspection_excel(
    save_path: str,
    batch_items: List[Dict[str, Any]],
    batch_summary: Optional[Dict[str, Any]] = None,
    progress_callback: Optional[Any] = None,
):
    """
    批量检测结果合并导出至单一 Excel 文件 (多 Sheet 结构):
    - Sheet 1: [批次检测总览] (包含宏观指标统计看板 + 各工件质检汇总清单 + 跳转对应 Sheet 的超链接)
    - Sheet 2..N: [各工件独立明细] (包含该工件的单板指标与该工件所有焊点的全量几何/质检明细)
    """
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()

    # 样式配置 (单例预定义，避免海量单元格重复构造)
    font_main_title = Font(name="微软雅黑", size=16, bold=True, color="1E3A8A")
    font_sheet_title = Font(name="微软雅黑", size=14, bold=True, color="1E3A8A")
    font_bold = Font(name="微软雅黑", size=11, bold=True)
    font_norm = Font(name="微软雅黑", size=10)
    font_header = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    font_link = Font(name="微软雅黑", size=10, underline="single", color="2563EB")

    fill_pass = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    fill_ng = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    fill_rev = PatternFill(start_color="CFFAFE", end_color="CFFAFE", fill_type="solid")
    fill_header = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    fill_card_head = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    style_pack = {
        "font_bold": font_bold,
        "font_norm": font_norm,
        "font_header": font_header,
        "fill_pass": fill_pass,
        "fill_ng": fill_ng,
        "fill_rev": fill_rev,
        "fill_header": fill_header,
        "thin_border": thin_border,
        "align_center": align_center,
        "align_left": align_left,
    }

    # 计算批次宏观统计
    total_boards = len(batch_items)
    pass_boards = sum(1 for item in batch_items if str(item.get("summary_stats", {}).get("board_status", "")).startswith("PASS"))
    ng_boards = total_boards - pass_boards
    pass_rate = (pass_boards / total_boards * 100.0) if total_boards > 0 else 0.0

    total_solders = sum(item.get("summary_stats", {}).get("total_solders", 0) for item in batch_items)
    total_void_ng = sum(item.get("summary_stats", {}).get("void_ng_count", 0) for item in batch_items)
    total_bridge = sum(item.get("summary_stats", {}).get("bridge_count", 0) for item in batch_items)
    total_insuff = sum(item.get("summary_stats", {}).get("insufficient_count", 0) for item in batch_items)
    total_cleared = sum(item.get("summary_stats", {}).get("manual_cleared_count", 0) for item in batch_items)
    total_elapsed = sum(item.get("summary_stats", {}).get("elapsed_ms", 0.0) for item in batch_items)
    avg_elapsed = (total_elapsed / total_boards) if total_boards > 0 else 0.0

    now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    if progress_callback:
        progress_callback(0, total_boards + 1, "正在构建批次总览看板...")

    # ==============================================================
    # 🌟 Sheet 1: 批次检测总览
    # ==============================================================
    ws_batch = wb.active
    ws_batch.title = "批次检测总览"

    # 主标题
    ws_batch["A1"] = "X射线 批量工业质检综合汇总报表"
    ws_batch["A1"].font = font_main_title
    ws_batch.merge_cells("A1:K1")

    # 宏观指标卡片
    overview_cards = [
        ("受检工件总数", total_boards, "批次最终判定", "PASS (全部合格)" if ng_boards == 0 else f"NG (含 {ng_boards} 片缺陷)"),
        ("合格工件数量", pass_boards, "批次合格率", f"{pass_rate:.1f}%"),
        ("不合格工件数", ng_boards, "报告导出时间", now_str),
        ("受检焊球总数", total_solders, "空洞超标总数", f"{total_void_ng} 处"),
        ("桥连短路总数", f"{total_bridge} 对", "虚焊少锡总数", f"{total_insuff} 处"),
        ("人工复核排除", f"{total_cleared} 处", "平均单片耗时", f"{avg_elapsed:.1f} ms"),
    ]

    for row_offset, (k1, v1, k2, v2) in enumerate(overview_cards, start=3):
        # 左侧键值
        c_k1 = ws_batch.cell(row=row_offset, column=1, value=k1)
        c_k1.font = font_bold
        c_k1.fill = fill_card_head
        c_k1.border = thin_border

        c_v1 = ws_batch.cell(row=row_offset, column=2, value=v1)
        c_v1.font = font_bold if k1 in ("受检工件总数", "合格工件数量") else font_norm
        c_v1.border = thin_border
        c_v1.alignment = align_center
        if k1 == "合格工件数量":
            c_v1.fill = fill_pass
        elif k1 == "不合格工件数" and ng_boards > 0:
            c_v1.fill = fill_ng

        # 右侧键值
        c_k2 = ws_batch.cell(row=row_offset, column=3, value=k2)
        c_k2.font = font_bold
        c_k2.fill = fill_card_head
        c_k2.border = thin_border

        c_v2 = ws_batch.cell(row=row_offset, column=4, value=v2)
        c_v2.font = font_bold if k2 in ("批次最终判定", "批次合格率") else font_norm
        c_v2.border = thin_border
        c_v2.alignment = align_center
        if k2 == "批次最终判定":
            c_v2.fill = fill_pass if ng_boards == 0 else fill_ng
            c_v2.font = Font(name="微软雅黑", size=11, bold=True, color="15803D" if ng_boards == 0 else "B91C1C")
        elif k2 == "批次合格率":
            c_v2.fill = fill_pass if pass_rate >= 90.0 else fill_ng

    # 批次工件汇总清单表
    batch_table_start_row = 10
    batch_headers = [
        "序号", "工件图像名称", "整板最终判定", "检测焊点数",
        "空洞超标", "桥连短路", "虚焊少锡", "人工排除",
        "最大空洞比", "最大少锡率", "检测用时(ms)", "详细数据工作表"
    ]

    for col_i, h in enumerate(batch_headers, start=1):
        cell = ws_batch.cell(row=batch_table_start_row, column=col_i, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border

    # 预先分配工件 Sheet 名称，确保唯一与双向超链接对应
    used_sheet_titles = {"批次检测总览"}
    sheet_title_list = []
    for idx, item in enumerate(batch_items, start=1):
        img_name = item.get("image_name") or os.path.basename(item.get("image_path", f"item_{idx}"))
        st = _sanitize_sheet_title(idx, img_name, used_sheet_titles)
        sheet_title_list.append(st)

    # 填充批次清单各行
    for idx, (item, sheet_title) in enumerate(zip(batch_items, sheet_title_list), start=1):
        r_row = batch_table_start_row + idx
        stats = item.get("summary_stats", {})
        img_name = item.get("image_name") or os.path.basename(item.get("image_path", f"item_{idx}"))
        b_status = stats.get("board_status", "UNKNOWN")
        is_pass = str(b_status).startswith("PASS")

        row_vals = [
            f"#{idx}",
            img_name,
            b_status,
            stats.get("total_solders", 0),
            stats.get("void_ng_count", 0),
            stats.get("bridge_count", 0),
            stats.get("insufficient_count", 0),
            stats.get("manual_cleared_count", 0),
            f"{float(stats.get('max_void_rate') or 0.0) * 100:.2f}%",
            f"-{float(stats.get('max_reduction_percent') or 0.0):.2f}%",
            f"{float(stats.get('elapsed_ms') or 0.0):.1f}",
            f"查看明细 ({sheet_title})",
        ]

        for col_i, val in enumerate(row_vals, start=1):
            cell = ws_batch.cell(row=r_row, column=col_i, value=val)
            cell.font = font_norm
            cell.border = thin_border
            cell.alignment = align_center

            if col_i == 2:
                # 工件名称居左
                cell.alignment = align_left
            elif col_i == 3:
                # 整板判定
                cell.font = Font(name="微软雅黑", size=10, bold=True, color="15803D" if is_pass else "B91C1C")
                cell.fill = fill_pass if is_pass else fill_ng
            elif col_i == 12:
                # 超链接到工件 Sheet
                cell.value = f'=HYPERLINK("#{sheet_title}!A1", "查看工件明细 ➜")'
                cell.font = font_link

    # ==============================================================
    # 🌟 Sheet 2..N: 各工件独立明细工作表
    # ==============================================================
    for idx, (item, sheet_title) in enumerate(zip(batch_items, sheet_title_list), start=1):
        img_name = item.get("image_name") or os.path.basename(item.get("image_path", f"item_{idx}"))
        if progress_callback:
            progress_callback(idx, total_boards + 1, f"正在生成工件明细表 ({idx}/{total_boards}): {img_name}")

        ws_item = wb.create_sheet(title=sheet_title)
        stats = item.get("summary_stats", {})
        records = item.get("records", [])
        b_status = stats.get("board_status", "UNKNOWN")
        is_pass = str(b_status).startswith("PASS")

        # 顶部标题栏
        ws_item["A1"] = f"X射线 工件质检明细单 - {img_name}"
        ws_item["A1"].font = font_sheet_title
        ws_item.merge_cells("A1:D1")

        # 返回总览快捷超链接
        ws_item["A2"] = '=HYPERLINK("#\'批次检测总览\'!A1", "◀ 返回批次检测总览")'
        ws_item["A2"].font = font_link
        ws_item.merge_cells("A2:B2")

        # 工件信息卡片 (第 4 行开始)
        item_summary_rows = [
            ("工件图片名称", img_name, "整板最终判定", b_status),
            ("检测焊点总数", stats.get("total_solders", 0), "空洞超标焊点", stats.get("void_ng_count", 0)),
            ("桥连短路对数", stats.get("bridge_count", 0), "虚焊少锡焊点", stats.get("insufficient_count", 0)),
            ("人工排除误检", stats.get("manual_cleared_count", 0), "检测用时 (ms)", f"{float(stats.get('elapsed_ms') or 0.0):.1f}"),
        ]

        for offset, (k1, v1, k2, v2) in enumerate(item_summary_rows, start=4):
            c_k1 = ws_item.cell(row=offset, column=1, value=k1)
            c_k1.font = font_bold
            c_k1.fill = fill_card_head
            c_k1.border = thin_border

            c_v1 = ws_item.cell(row=offset, column=2, value=v1)
            c_v1.font = font_norm
            c_v1.border = thin_border
            c_v1.alignment = align_center

            c_k2 = ws_item.cell(row=offset, column=3, value=k2)
            c_k2.font = font_bold
            c_k2.fill = fill_card_head
            c_k2.border = thin_border

            c_v2 = ws_item.cell(row=offset, column=4, value=v2)
            c_v2.font = font_norm
            c_v2.border = thin_border
            c_v2.alignment = align_center

            if k2 == "整板最终判定":
                c_v2.font = Font(name="微软雅黑", size=11, bold=True, color="15803D" if is_pass else "B91C1C")
                c_v2.fill = fill_pass if is_pass else fill_ng

        # 焊点全量明细表 (从第 9 行开始)
        _append_solder_records_table(ws_item, records, start_row=9, styles=style_pack)
        # 极速预置列宽与网格线 (免除数万单元格重复扫描)
        _apply_detail_sheet_col_widths(ws_item)

    # 批次总览自适应宽度
    _apply_sheet_styles_and_col_width(ws_batch, get_column_letter)

    if progress_callback:
        progress_callback(total_boards, total_boards + 1, "正在保存 Excel 表格...")

    wb.save(save_path)
