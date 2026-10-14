# -*- coding: utf-8 -*-
"""
X射线 Word 模板报告生成器 (Word Reporter)
---------------------------------------
基于用户可视化的 Word 模板 (.docx) 生成工业质检报告：
1. 支持用户自由在 Word 中修改排版、标题、LOGO、表格样式与文字；
2. 稳健的段落级占位符替换 (解决 Word 编辑后 run 分词打碎的问题)；
3. 智能替换全景缺陷标注图 (等比缩放保真居中)；
4. 动态扩展缺陷与复核明细表 (支持高亮底色与状态标识)；
5. 支持输出高质量 .docx，并可选通过 Word/WPS 内核一键转换为 .pdf。
"""
from __future__ import annotations

import datetime
import os
import re
from typing import Any, Dict, List, Optional
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from PIL import Image as PILImage

from desktop_app.core.inspection_manager import SolderBallRecord
from desktop_app.utils.defect_renderer import render_defect_only_preview


def set_cell_background(cell, hex_color: str):
    """设置单元格背景颜色"""
    hex_color = hex_color.lstrip('#')
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def set_cell_margins(cell, top=80, bottom=80, left=80, right=80):
    """设置单元格内边距"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def _replace_paragraph_placeholders(paragraph, replacements: Dict[str, str]):
    """
    稳健地替换段落中的 {{KEY}} 占位符。
    解决 Word 在编辑保存后常将单个占位符打碎在不同 run 节点的问题。
    """
    full_text = "".join(r.text for r in paragraph.runs)
    matched = False
    for k in replacements:
        if k in full_text:
            matched = True
            break
    if not matched:
        return

    # 进行全局替换
    new_text = full_text
    for k, v in replacements.items():
        new_text = new_text.replace(k, str(v))

    # 保留第一个 run 的基本字体属性，并将新文本写入
    if paragraph.runs:
        font_name = paragraph.runs[0].font.name
        font_size = paragraph.runs[0].font.size
        bold = paragraph.runs[0].font.bold
        color = paragraph.runs[0].font.color.rgb if paragraph.runs[0].font.color else None

        paragraph.runs[0].text = new_text
        for r in paragraph.runs[1:]:
            r.text = ""

        # 尝试保留中文字体
        try:
            if font_name:
                paragraph.runs[0]._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
        except Exception:
            pass


def _replace_table_placeholders(table, replacements: Dict[str, str]):
    """遍历表格所有单元格，替换普通占位符"""
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                _replace_paragraph_placeholders(p, replacements)


def convert_docx_to_pdf(docx_path: str, pdf_path: str) -> bool:
    """使用 Windows 本地 WPS / Word (COM 组件) 将 docx 转为高保真 PDF"""
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    for app_name in ["kwps.Application", "wps.Application", "Word.Application"]:
        try:
            app = win32com.client.Dispatch(app_name)
            app.Visible = False
            app.DisplayAlerts = 0
            doc = app.Documents.Open(os.path.abspath(docx_path))
            if hasattr(doc, "ExportAsFixedFormat"):
                doc.ExportAsFixedFormat(os.path.abspath(pdf_path), 17)
            else:
                doc.SaveAs(os.path.abspath(pdf_path), FileFormat=17)
            doc.Close(SaveChanges=0)
            app.Quit()
            pythoncom.CoUninitialize()
            return True
        except Exception:
            continue
    pythoncom.CoUninitialize()
    return False


def generate_word_report(
    template_path: str,
    output_docx_path: str,
    image_name: str,
    summary_stats: Dict[str, Any],
    records: List[SolderBallRecord],
    preview_img_path: Optional[str] = None,
    image_path: Optional[str] = None,
    bridge_defects: Optional[List[Dict[str, Any]]] = None,
    inspect_mode: str = "全项综合检测 (空洞+连锡+虚焊)",
    output_pdf_path: Optional[str] = None,
) -> str:
    """
    基于 Word 模板生成质检报告。
    
    参数:
    - template_path: Word 模板路径 (.docx)
    - output_docx_path: 输出的 Word 报告路径 (.docx)
    - image_name: 待测工件名称
    - summary_stats: 质检汇总指标字典
    - records: 焊点检测记录列表
    - preview_img_path: 高清预览图路径 (若无会自动渲染纯缺陷图)
    - image_path: 原始图像路径
    - bridge_defects: 桥连缺陷列表
    - inspect_mode: 检测模式说明
    - output_pdf_path: 若指定，则尝试同步转换为 PDF 文件
    
    返回:
    - 成功生成的 docx 报告绝对路径
    """
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"找不到指定的 Word 模板文件: {template_path}")

    # 1. 准备统计指标
    board_status = summary_stats.get("board_status", "UNKNOWN")
    now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    elapsed_ms = f"{float(summary_stats.get('elapsed_ms') or 0.0):.1f}"
    total_solders = str(summary_stats.get("total_solders", 0))
    v_ng = summary_stats.get("void_ng_count", 0)
    b_ng = summary_stats.get("bridge_count", 0)
    i_ng = summary_stats.get("insufficient_count", 0)
    c_rev = summary_stats.get("manual_cleared_count", 0)
    max_void = f"{float(summary_stats.get('max_void_rate') or 0.0) * 100:.1f}%"

    replacements = {
        "{{IMAGE_NAME}}": image_name,
        "{{INSPECT_MODE}}": inspect_mode,
        "{{INSPECT_TIME}}": now_str,
        "{{ELAPSED_MS}}": elapsed_ms,
        "{{BOARD_STATUS}}": board_status,
        "{{TOTAL_SOLDERS}}": total_solders,
        "{{VOID_NG_COUNT}}": f"{v_ng} 处" if v_ng > 0 else "0",
        "{{BRIDGE_COUNT}}": f"{b_ng} 处" if b_ng > 0 else "0",
        "{{INSUFFICIENT_COUNT}}": f"{i_ng} 处" if i_ng > 0 else "0",
        "{{MANUAL_CLEARED_COUNT}}": f"{c_rev} 处" if c_rev > 0 else "0",
        "{{MAX_VOID_RATE}}": max_void,
    }

    # 2. 准备纯缺陷高清标注图
    actual_preview = preview_img_path
    target_img = image_path or preview_img_path
    if target_img and os.path.exists(target_img):
        try:
            actual_preview = render_defect_only_preview(
                image_path=target_img,
                records=records,
                bridge_defects=bridge_defects or [],
            )
        except Exception as e:
            print(f"[WARN] 渲染单图纯缺陷标注图失败: {e}")

    # 3. 加载模板文档
    doc = Document(template_path)

    # 替换所有主文档正文段落
    for p in doc.paragraphs:
        _replace_paragraph_placeholders(p, replacements)

    # 4. 遍历所有表格替换占位符，并处理整板判定颜色
    for tbl in doc.tables:
        _replace_table_placeholders(tbl, replacements)

        # 检查是否包含判定单元格，自适应设置颜色
        for row in tbl.rows:
            for cell in row.cells:
                if board_status in cell.text:
                    if board_status.startswith("PASS"):
                        set_cell_background(cell, "DCFCE7")  # 浅绿
                    elif "人工复核" in board_status:
                        set_cell_background(cell, "CFFAFE")  # 浅青
                    else:
                        set_cell_background(cell, "FEE2E2")  # 浅红

    # 5. 替换图片占位符 {{PREVIEW_IMAGE}}
    image_inserted = False
    for p in doc.paragraphs:
        if "{{PREVIEW_IMAGE}}" in p.text:
            p.text = ""  # 清空文字
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if actual_preview and os.path.exists(actual_preview):
                try:
                    # 计算等比尺寸：最大宽 16cm，最大高 9.5cm
                    with PILImage.open(actual_preview) as pil_im:
                        pw, ph = pil_im.size
                    aspect = float(pw) / float(ph) if ph > 0 else 1.5

                    max_w_cm = 16.0
                    max_h_cm = 9.5
                    if aspect >= (max_w_cm / max_h_cm):
                        render_w = Cm(max_w_cm)
                        render_h = Cm(max_w_cm / aspect)
                    else:
                        render_h = Cm(max_h_cm)
                        render_w = Cm(max_h_cm * aspect)

                    run = p.add_run()
                    run.add_picture(actual_preview, width=render_w, height=render_h)
                    image_inserted = True
                except Exception as e:
                    p.text = f"[图片加载异常: {e}]"
            else:
                p.text = "[暂无标注预览图像]"
            break

    # 6. 填充缺陷与复核焊点明细表
    defect_records = [r for r in records if r.is_ng or r.is_manually_reviewed]

    for tbl in doc.tables:
        # 寻找包含 {{DEFECT_ROW}} 的表格
        target_row_idx = -1
        for row_i, row in enumerate(tbl.rows):
            if any("{{DEFECT_ROW}}" in cell.text for cell in row.cells):
                target_row_idx = row_i
                break

        if target_row_idx != -1:
            # 找到模板行
            template_row = tbl.rows[target_row_idx]

            if not defect_records:
                # 全板合格无缺陷
                for idx, cell in enumerate(template_row.cells):
                    if idx == 0:
                        cell.text = "整板未检出任何空洞、连锡或虚焊缺陷，质量合格。"
                        set_cell_background(cell, "DCFCE7")
                        p = cell.paragraphs[0]
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        if p.runs:
                            p.runs[0].font.bold = True
                            p.runs[0].font.color.rgb = RGBColor(0x16, 0xA3, 0x4A)
                    else:
                        cell.text = "-"
                        set_cell_background(cell, "DCFCE7")
            else:
                # 遍历缺陷记录，动态添加行
                max_show = 150
                records_to_show = defect_records[:max_show]

                for r in records_to_show:
                    is_rev = (r.is_manually_reviewed and r.manual_status == "PASS")
                    concl = "已排除误检(PASS)" if is_rev else "NG 不良"
                    rcx = float(r.cx) if r.cx is not None else 0.0
                    rcy = float(r.cy) if r.cy is not None else 0.0
                    r_void = float(r.void_rate) if r.void_rate is not None else 0.0
                    r_reduct = float(r.reduction_percent) if r.reduction_percent is not None else 0.0

                    row_vals = [
                        f"#{r.index}",
                        f"({rcx:.1f}, {rcy:.1f})",
                        f"{r_void * 100:.1f}%",
                        f"-{r_reduct:.1f}%" if r_reduct > 0 else "0.0%",
                        "桥连短路NG" if r.is_bridge else "正常",
                        ", ".join(r.ng_reasons) if r.ng_reasons else "异常",
                        concl,
                    ]

                    # 插入新行
                    new_row = tbl.add_row()
                    bg_color = "CFFAFE" if is_rev else "FFF1F2"

                    for col_idx, val in enumerate(row_vals):
                        if col_idx < len(new_row.cells):
                            cell = new_row.cells[col_idx]
                            cell.text = val
                            set_cell_background(cell, bg_color)
                            set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
                            p = cell.paragraphs[0]
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            p.paragraph_format.space_after = Pt(0)
                            if p.runs:
                                p.runs[0].font.size = Pt(8)
                                p.runs[0].font.name = '微软雅黑'
                                try:
                                    p.runs[0]._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
                                except Exception:
                                    pass

                # 删除模板占位行
                tr = template_row._tr
                tbl._tbl.remove(tr)

    # 7. 保存最终 Word 报告
    os.makedirs(os.path.dirname(os.path.abspath(output_docx_path)), exist_ok=True)
    doc.save(output_docx_path)
    print(f"[OK] Word 报告已生成: {output_docx_path}")

    # 8. 若要求转为 PDF 且指定了路径
    if output_pdf_path:
        os.makedirs(os.path.dirname(os.path.abspath(output_pdf_path)), exist_ok=True)
        pdf_ok = convert_docx_to_pdf(output_docx_path, output_pdf_path)
        if pdf_ok:
            print(f"[OK] Word 报告已同步转换为 PDF: {output_pdf_path}")
        else:
            print(f"[INFO] 未能自动调用本地 Word COM 转为 PDF，请使用 Word/WPS 打开并另存为 PDF")

    return os.path.abspath(output_docx_path)
